"""
Convert the KT markdown documents into Word (.docx) files.

Usage (from the KT folder):
    ..\\backend\\venv\\Scripts\\python.exe convert-to-docx.py

Renders: headings, paragraphs, bullet/numbered lists, blockquotes,
pipe tables, fenced code blocks (whitespace preserved so the ASCII
architecture diagrams stay aligned), horizontal rules, and the inline
**bold** / `code` / [link](url) forms used in these documents.
"""

import os
import re
import glob

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

CODE_FONT = "Consolas"
CODE_SIZE = Pt(8.5)
CODE_SHADE = "F2F2F2"
LINK_COLOR = RGBColor(0x0B, 0x5C, 0xAB)
CODE_COLOR = RGBColor(0xA3, 0x1D, 0x1D)

INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*"          # bold
    r"|`[^`]+`"                # inline code
    r"|\[[^\]]+\]\([^)]+\))"   # link
)


# ----------------------------------------------------------------- helpers

def shade(element, fill):
    """Apply a background fill to a paragraph or table cell."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFBFBF")
    borders.append(bottom)
    pPr.append(borders)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def add_inline(paragraph, text):
    """Write text into a paragraph, honouring inline markdown."""
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(9)
            run.font.color.rgb = CODE_COLOR
        elif token.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if m:
                run = paragraph.add_run(m.group(1))
                run.font.color.rgb = LINK_COLOR
                run.underline = True
            else:
                paragraph.add_run(token)
        else:
            paragraph.add_run(token)


def add_code_block(doc, lines):
    """One shaded, monospaced paragraph; newlines kept so diagrams align."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.12)
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    shade(p._p.get_or_add_pPr(), CODE_SHADE)

    for i, line in enumerate(lines):
        run = p.add_run()
        if i:
            run.add_break()
        run.add_text(line)
        run.font.name = CODE_FONT
        run.font.size = CODE_SIZE
        # ensure the monospace font also applies to non-ASCII diagram glyphs
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), CODE_FONT)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def add_table(doc, rows):
    header = split_row(rows[0])
    body = [split_row(r) for r in rows[2:]]
    ncols = max([len(header)] + [len(r) for r in body]) if body else len(header)

    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i in range(ncols):
        cell = table.rows[0].cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        add_inline(para, header[i] if i < len(header) else "")
        for run in para.runs:
            run.bold = True
        shade(cell._tc.get_or_add_tcPr(), "DEEAF6")

    for r in body:
        cells = table.add_row().cells
        for i in range(ncols):
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, r[i] if i < len(r) else "")

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ------------------------------------------------------------- conversion

def convert(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as fh:
        lines = fh.read().replace("\r\n", "\n").split("\n")

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    for section in doc.sections:
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)

    i = 0
    first_h1 = True
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            while block and not block[-1].strip():
                block.pop()
            if block:
                add_code_block(doc, block)
            continue

        # pipe table (needs a |---| separator on the next line)
        if (
            stripped.startswith("|")
            and i + 1 < n
            and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1])
        ):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue

        # horizontal rule
        if re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", line):
            add_hr(doc)
            i += 1
            continue

        # blank line
        if not stripped:
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and first_h1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(re.sub(r"[*`]", "", text))
                run.bold = True
                run.font.size = Pt(20)
                run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
                p.paragraph_format.space_after = Pt(10)
                first_h1 = False
            else:
                h = doc.add_heading(level=min(level, 4))
                for run in list(h.runs):
                    run.text = ""
                add_inline(h, text)
                h.paragraph_format.space_before = Pt(10)
                h.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # blockquote (may span several lines)
        if stripped.startswith(">"):
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.25)
            pf.space_before = Pt(4)
            pf.space_after = Pt(6)
            shade(p._p.get_or_add_pPr(), "FFF7E0")
            add_inline(p, " ".join(q for q in quote if q))
            for run in p.runs:
                run.italic = True
            continue

        # bullet list item
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if m:
            depth = len(m.group(1)) // 2
            style = "List Bullet" if depth == 0 else "List Bullet 2"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2).strip())
            i += 1
            continue

        # numbered list item
        m = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if m:
            depth = len(m.group(1)) // 2
            style = "List Number" if depth == 0 else "List Number 2"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2).strip())
            i += 1
            continue

        # plain paragraph
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(docx_path)


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    md_files = sorted(glob.glob(os.path.join(here, "*.md")))
    if not md_files:
        print("No markdown files found in", here)
        return
    for md in md_files:
        out = os.path.splitext(md)[0] + ".docx"
        convert(md, out)
        size = os.path.getsize(out) / 1024.0
        print("OK  {0:<46} -> {1} ({2:.0f} KB)".format(
            os.path.basename(md), os.path.basename(out), size))
    print("\nConverted {0} file(s).".format(len(md_files)))


if __name__ == "__main__":
    main()
